"""
zotero_item_type_fix.py

Fixes Zotero library items that were imported as "Journal Article" but are
actually Newspaper or Magazine articles per the source database's
reference_type.

Background: BibTeX has no @newspaper or @magazine entry type -- @article is
the only periodical entry type the format defines. db2bibtex.exporter's
ENTRY_TYPE_MAP therefore maps "Journal Article", "Magazine Article", and
"Newspaper Article" reference_type values all to BibTeX's "article" type
(see exporter.py), and Zotero's BibTeX-import translator unconditionally
turns every @article entry into item type "journalArticle" on import,
regardless of what's in the journal/publisher field. That's a real
limitation of the BibTeX round trip, not a bug in the reference_type
mapping -- confirmed against live data (e.g. Seattle Times items authored
by Mapes have reference_type = 'Newspaper Article' in the source DB and are
correctly mapped to "article" on export, but land in Zotero as "Journal
Article" because that's the only place @article can go).

This module is the fix: it scans an already-imported Zotero library,
figures out which "Journal Article" items are really newspaper/magazine
articles by looking their publication_id back up in the source database,
and converts just those items' Zotero item type in place -- carrying over
every field the new type still supports and backfilling place/section from
columns (place_published, news_section) that db2bibtex's BibTeX export
doesn't currently carry through at all.

Standalone feature: does not import db2bibtex.exporter's BibTeX-building
functions and never writes to the SQL Server database -- it only reads
from it (via an external, gitignored lookup query, same convention as
db2bibtex.exporter.load_query) to decide what each item's real type should
be. Reuses db2bibtex.compare's publication_id-extraction regex, since it's
the same join key compare.py already established as the one reliable
identifier across a library containing items imported under several
historical exporter versions (unlike the `keywords` tag, which was added
later and isn't present on every item).

Design notes:
- Matching a Zotero item back to its source-DB row is done via the
  `publication_id` embedded in a "Source DB: publication_id X; pub_number
  Y; catalog_id Z" child note -- confirmed against a live library sample:
  Zotero's BibTeX-import translator turns the exporter's `note` field into
  a standalone child note item, not the parent's `extra` field (which
  import leaves empty). find_pub_ids_by_parent_note() does one bulk,
  paginated scan of all 'note'-type items (rather than a child-items call
  per candidate) and extracts the id with db2bibtex.compare.extract_pub_id()
  -- the same regex compare.py uses on raw .bib text, just applied to the
  note's HTML instead. NOT the `keywords` tag -- see compare.py's module
  docstring for why that's unreliable across a mixed-vintage library.
- Real table/column names (LTERMETA.dbo.publication et al.) are kept out
  of this tracked module the same way exporter.py keeps them out of
  query.sql's caller: the lookup query text lives in an external,
  gitignored query_item_type_lookup.sql file (matches the `query_*.sql` /
  `!query_*.sql.example` .gitignore pattern already used for
  query_not_hja_lter.sql), read via exporter.load_query(). See
  query_item_type_lookup.sql.example for the expected shape -- it must
  contain exactly one `{ph}` placeholder, which this module fills with the
  right number of `?` parameter markers for however many publication_ids
  it needs to look up in one round trip.
- Item-type conversion is a full field remap built from a fresh
  zot.item_template(new_type), not a patch to itemType in place: Zotero's
  own "Change Item Type" UI drops any field with no home in the new type's
  schema (e.g. volume/issue when converting to newspaperArticle) and
  leaves newly-available fields (place, section) at their template
  default, so build_new_data() does the same instead of mutating itemType
  and hoping stale fields are silently ignored.
- process_item() re-fetches the item immediately before writing, same
  reasoning as zotero_author_complete.process_item: guards against
  clobbering an edit made elsewhere since find_candidate_items() ran.
"""

from __future__ import annotations

import configparser
import logging
from dataclasses import dataclass, field
from pathlib import Path

try:
    import pyodbc
except ImportError:  # pragma: no cover -- exercised via PyodbcMissingError tests
    pyodbc = None

try:
    from pyzotero import zotero
except ImportError:  # pragma: no cover -- exercised via ZoteroDepsMissingError tests
    zotero = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:  # pragma: no cover -- exercised via ZoteroDepsMissingError tests
    Workbook = None
    Font = None

from db2bibtex.compare import extract_pub_id
from db2bibtex.db import build_conn_str
from db2bibtex.exporter import DEFAULT_DRIVER, PyodbcMissingError, QueryFileMissingError, load_query

log = logging.getLogger(__name__)

# reference_type substring (lowercased) -> Zotero item type. Only reference
# types that BibTeX's @article limitation can misclassify need an entry
# here -- anything else db2bibtex.exporter.guess_entry_type() already routes
# to its own distinct BibTeX/Zotero type (book, incollection, ...), so this
# tool never needs to touch it.
REFERENCE_TYPE_TO_ZOTERO_TYPE = {
    "newspaper article": "newspaperArticle",
    "magazine article": "magazineArticle",
}

DEFAULT_QUERY_FILE = "query_item_type_lookup.sql"


class ZoteroDepsMissingError(RuntimeError):
    """Raised when a Zotero operation is attempted but pyzotero/openpyxl
    are not importable."""


@dataclass
class Config:
    library_id: str
    library_type: str  # 'group' or 'user'
    api_key: str
    server: str
    database: str
    driver: str = DEFAULT_DRIVER
    uid: "str | None" = None
    pwd: "str | None" = None
    trusted_connection: bool = False
    trust_server_certificate: bool = True
    encrypt: bool = True
    query_file: str = DEFAULT_QUERY_FILE
    dry_run: bool = True
    audit_path: str = "zotero_item_type_fix_audit.xlsx"
    collection_key: "str | None" = None


@dataclass
class AuditRow:
    item_key: str
    title: str
    old_item_type: str
    new_item_type: str
    publication_id: str
    status: str  # 'updated' | 'dry-run' | 'skipped-no-pub-id' | 'skipped-not-in-db' |
    # 'skipped-no-change-needed' | 'error'
    detail: str = ""


@dataclass
class RunResult:
    rows: list = field(default_factory=list)
    counts: dict = field(default_factory=dict)


def load_config(config_path: "str | Path") -> dict:
    """Read [Zotero]/[Database]/[Driver] settings from a db_config.ini-style
    file (see db_config.ini.example -- the same file db2bibtex.exporter and
    db2bibtex.zotero_author_complete read their own sections from).

    Returns:
        A dict with keys: library_id, library_type, api_key, server,
        database, uid, pwd, driver, trust_server_certificate.
    """
    cp = configparser.ConfigParser()
    read_files = cp.read(config_path)
    if not read_files:
        raise FileNotFoundError(f"Could not read config file: {config_path}")

    zsec = cp["Zotero"] if cp.has_section("Zotero") else {}
    dsec = cp["Database"] if cp.has_section("Database") else {}
    drv = cp["Driver"] if cp.has_section("Driver") else {}

    return {
        "library_id": zsec.get("library_id"),
        "library_type": zsec.get("library_type"),
        "api_key": zsec.get("api_key"),
        "server": dsec.get("server"),
        "database": dsec.get("database"),
        "uid": dsec.get("username"),
        "pwd": dsec.get("password"),
        "driver": drv.get("driver", DEFAULT_DRIVER).strip("{}") if drv else DEFAULT_DRIVER,
        "trust_server_certificate": cp.getboolean(
            "Database", "trustservercertificate", fallback=True
        )
        if cp.has_section("Database")
        else True,
    }


# ---------------------------------------------------------------------------
# Source-database lookup
# ---------------------------------------------------------------------------

def _require_pyodbc() -> None:
    if pyodbc is None:
        raise PyodbcMissingError(
            "pyodbc is required for the source-database lookup but is not "
            "importable (install it and/or the unixODBC system library)"
        )


def _build_conn_str(cfg: Config) -> str:
    return build_conn_str(
        cfg.driver,
        cfg.server,
        cfg.database,
        cfg.trusted_connection,
        cfg.uid,
        cfg.pwd,
        cfg.encrypt,
        cfg.trust_server_certificate,
    )


def fetch_reference_types(cfg: Config, publication_ids: "set[str]") -> dict:
    """Look up reference_type/pub_type/place_published/news_section for a
    batch of publication_ids in one round trip.

    Args:
        cfg: Config with database connection settings and query_file.
        publication_ids: publication_id strings to look up (as extracted
            from Zotero items' linked notes via find_pub_ids_by_parent_note).

    Returns:
        Dict of publication_id (str) -> {"reference_type", "pub_type",
        "place_published", "news_section"}. IDs with no matching row are
        simply absent from the result.
    """
    if not publication_ids:
        return {}
    _require_pyodbc()

    query_template = load_query(cfg.query_file)
    ids = sorted(publication_ids, key=int)
    placeholders = ",".join("?" * len(ids))
    query = query_template.format(ph=placeholders)

    conn = pyodbc.connect(_build_conn_str(cfg))
    try:
        cursor = conn.cursor()
        cursor.execute(query, [int(pid) for pid in ids])
        rows = {}
        for row in cursor.fetchall():
            rows[str(row.publication_id)] = {
                "reference_type": row.reference_type,
                "pub_type": row.pub_type,
                "place_published": row.place_published,
                "news_section": row.news_section,
            }
        return rows
    finally:
        conn.close()


def guess_target_item_type(reference_type, pub_type=None) -> "str | None":
    """Map a source-DB reference_type/pub_type to the Zotero item type it
    should have been imported as, or None if no rule applies (leave the
    item alone -- covers real journal articles and anything this tool
    doesn't know how to fix).

    Args:
        reference_type: The publication's reference_type column.
        pub_type: The publication's pub_type column (fallback, matching
            db2bibtex.exporter.guess_entry_type's own fallback order).

    Returns:
        A Zotero item type string ("newspaperArticle", "magazineArticle"),
        or None.
    """
    rt = (reference_type or pub_type or "").strip().lower()
    for key, zotero_type in REFERENCE_TYPE_TO_ZOTERO_TYPE.items():
        if key in rt:
            return zotero_type
    return None


# ---------------------------------------------------------------------------
# Zotero item-type conversion
# ---------------------------------------------------------------------------

def _require_zotero() -> None:
    if zotero is None:
        raise ZoteroDepsMissingError(
            "pyzotero is required but is not importable. "
            "Install it with: pip install -e '.[zotero]'"
        )


def get_client(cfg: Config):
    """Build a pyzotero client from a Config."""
    _require_zotero()
    return zotero.Zotero(cfg.library_id, cfg.library_type, cfg.api_key)


def find_candidate_items(zot, collection_key: "str | None") -> list:
    """Paginate the full library (or one collection) and return every item
    currently typed "journalArticle" -- the only type BibTeX's @article
    limitation can produce incorrectly, so it's the full candidate pool
    this tool ever needs to re-check against the source database."""
    if collection_key:
        items = zot.everything(zot.collection_items(collection_key))
    else:
        items = zot.everything(zot.top())
    return [it for it in items if it["data"].get("itemType") == "journalArticle"]


def find_pub_ids_by_parent_note(zot, collection_key: "str | None") -> dict:
    """One paginated scan of 'note' child items (library-wide, or scoped to
    one collection), mapping parentItem key -> publication_id.

    Confirmed against a live library sample: Zotero's BibTeX-import
    translator turns the exporter's `note` field (always populated --
    build_bibtex_entry() unconditionally writes "Source DB: publication_id
    X; pub_number Y; catalog_id Z" into it) into a standalone child note
    item ("<p>Notes: ...; Source DB: publication_id 5378; ...</p>"), *not*
    the parent item's `extra` field, which import leaves empty. This is a
    live-API-specific finding, distinct from db2bibtex.compare's matching
    (which parses the same "Source DB: publication_id" text, but out of
    raw .bib file text -- there, it really is the `note`/`annote` field on
    the entry itself, since compare.py never touches the Zotero API).
    extract_pub_id() is the same regex either way; only where it's applied
    differs.
    """
    if collection_key:
        notes = zot.everything(zot.collection_items(collection_key, itemType="note"))
    else:
        notes = zot.everything(zot.items(itemType="note"))
    result = {}
    for note in notes:
        data = note["data"]
        parent = data.get("parentItem")
        if not parent:
            continue
        pub_id = extract_pub_id(data.get("note", "") or "")
        if pub_id:
            result[parent] = pub_id
    return result


def build_new_data(old_data: dict, template: dict, db_row: dict) -> dict:
    """Build the new item `data` dict for converting old_data to the type
    template came from.

    template is a fresh skeleton from zot.item_template(new_type) -- its
    keys are exactly the fields Zotero considers valid for that type. Any
    field present in old_data but absent from template has no home in the
    new type and is dropped, matching what Zotero's own "Change Item Type"
    UI does; any shared field (title, creators, date, url, tags,
    collections, relations, ...) carries over unchanged.

    Args:
        old_data: The item's current `data` dict (from a freshly re-fetched
            item, so `key`/`version` are current).
        template: zot.item_template(new_type)'s return value.
        db_row: The matching row from fetch_reference_types(), used to
            backfill place/section on newspaperArticle -- data BibTeX
            export doesn't currently carry through at all.

    Returns:
        A new `data` dict with `key`/`version` preserved from old_data.
    """
    new_data = dict(template)
    for field_name, value in old_data.items():
        # itemType must stay whatever `template` says (the new type) -- it's
        # a key in both old_data and template, so the generic "shared field"
        # copy below would otherwise silently overwrite it back to the old
        # type and defeat the whole conversion.
        if field_name == "itemType":
            continue
        if field_name in new_data:
            new_data[field_name] = value
    new_data["key"] = old_data["key"]
    if "version" in old_data:
        new_data["version"] = old_data["version"]

    if new_data.get("itemType") == "newspaperArticle":
        place = (db_row or {}).get("place_published")
        if place and not new_data.get("place"):
            new_data["place"] = place
        section = (db_row or {}).get("news_section")
        if section and not new_data.get("section"):
            new_data["section"] = section

    return new_data


def process_item(zot, item: dict, publication_id: "str | None", db_row: "dict | None", cfg: Config) -> AuditRow:
    """Decide and (unless dry_run) apply the item-type fix for one item."""
    data = item["data"]
    key = data["key"]
    title = data.get("title", "")[:80]
    old_type = data.get("itemType", "")

    if not publication_id:
        log.info("skip (no publication_id in a linked note): %s [%s]", title, key)
        return AuditRow(key, title, old_type, "", "", "skipped-no-pub-id")

    if db_row is None:
        log.info("skip (publication_id %s not found in database): %s [%s]", publication_id, title, key)
        return AuditRow(key, title, old_type, "", publication_id, "skipped-not-in-db")

    reference_type = db_row.get("reference_type") or db_row.get("pub_type") or ""
    new_type = guess_target_item_type(db_row.get("reference_type"), db_row.get("pub_type"))
    detail = f"reference_type={reference_type.strip()!r}"

    if not new_type or new_type == old_type:
        return AuditRow(key, title, old_type, new_type or old_type, publication_id, "skipped-no-change-needed", detail)

    if cfg.dry_run:
        log.info("would convert: %s [%s]: %s -> %s", title, key, old_type, new_type)
        return AuditRow(key, title, old_type, new_type, publication_id, "dry-run", detail)

    # Re-fetch immediately before writing so we hold the current version
    # (guards against edits made elsewhere since find_candidate_items ran).
    fresh_item = zot.item(key)
    template = zot.item_template(new_type)
    fresh_item["data"] = build_new_data(fresh_item["data"], template, db_row)
    try:
        zot.update_item(fresh_item)
        log.info("converted: %s [%s]: %s -> %s", title, key, old_type, new_type)
        return AuditRow(key, title, old_type, new_type, publication_id, "updated", detail)
    except Exception as exc:  # noqa: BLE001 -- surfaced via audit workbook, not swallowed
        log.error("update failed: %s [%s]: %s", title, key, exc)
        return AuditRow(key, title, old_type, new_type, publication_id, "error", str(exc))


AUDIT_HEADERS = [
    "item_key",
    "title",
    "old_item_type",
    "new_item_type",
    "publication_id",
    "status",
    "detail",
]
AUDIT_COLUMN_WIDTHS = {"A": 12, "B": 45, "C": 18, "D": 18, "E": 14, "F": 24, "G": 40}


def _require_openpyxl() -> None:
    if Workbook is None:
        raise ZoteroDepsMissingError(
            "openpyxl is required to write the audit workbook but is not "
            "importable. Install it with: pip install -e '.[zotero]'"
        )


def write_audit_xlsx(rows: list, path: str) -> None:
    """Write the audit trail as a real .xlsx workbook (native Unicode, real
    cell boundaries -- same rationale as zotero_author_complete's audit
    workbook)."""
    _require_openpyxl()
    wb = Workbook()
    ws = wb.active
    ws.title = "audit"
    ws.append(AUDIT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append(
            [r.item_key, r.title, r.old_item_type, r.new_item_type, r.publication_id, r.status, r.detail]
        )
    for col_letter, width in AUDIT_COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width
    wb.save(path)


def run(cfg: Config, progress_callback=print) -> RunResult:
    """Scan the library (or one collection), fix every candidate item, and
    write the audit workbook. Used by both the CLI and tests."""
    zot = get_client(cfg)

    progress_callback(
        "Scanning library for Journal Article items"
        + (f" in collection {cfg.collection_key}" if cfg.collection_key else "")
        + "..."
    )
    candidates = find_candidate_items(zot, cfg.collection_key)
    progress_callback(f"Found {len(candidates)} Journal Article item(s) to check")

    progress_callback(
        "Scanning all notes for Source DB publication_ids"
        + ("..." if cfg.collection_key else " -- this is a full-library scan and can take "
           "several minutes on a large library; use --collection to test faster first...")
    )
    pub_ids_by_key = find_pub_ids_by_parent_note(zot, cfg.collection_key)
    progress_callback(f"Found {len(pub_ids_by_key)} linked note(s) with a Source DB publication_id")

    db_rows = fetch_reference_types(
        cfg, {pub_ids_by_key[item["data"]["key"]] for item in candidates if item["data"]["key"] in pub_ids_by_key}
    )
    progress_callback(f"Matched {len(db_rows)} of {len(pub_ids_by_key)} publication_id(s) in the source database")

    rows = [
        process_item(zot, item, pub_ids_by_key.get(item["data"]["key"]), db_rows.get(pub_ids_by_key.get(item["data"]["key"])), cfg)
        for item in candidates
    ]

    write_audit_xlsx(rows, cfg.audit_path)
    counts: dict = {}
    for r in rows:
        counts[r.status] = counts.get(r.status, 0) + 1
    progress_callback(f"Done. {counts}")
    if cfg.dry_run:
        progress_callback(
            f"Dry-run only -- rerun with --live to apply changes. Review {cfg.audit_path} first."
        )
    return RunResult(rows=rows, counts=counts)
