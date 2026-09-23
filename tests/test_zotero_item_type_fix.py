"""Tests for db2bibtex.zotero_item_type_fix: reference_type -> Zotero item
type mapping, the item-type field remap, and the full fix pipeline.

No real Zotero/SQL Server calls are made -- pyzotero and pyodbc are mocked
throughout, same convention as test_exporter.py / test_zotero_author_complete.py.
"""

from __future__ import annotations

import configparser
from unittest.mock import MagicMock

import pytest

from db2bibtex import zotero_item_type_fix as fix


# ---------------------------------------------------------------------------
# guess_target_item_type
# ---------------------------------------------------------------------------

def test_guess_target_item_type_newspaper():
    assert fix.guess_target_item_type("Newspaper Article", None) == "newspaperArticle"


def test_guess_target_item_type_newspaper_with_padding_and_case():
    # LTERMETA's reference_type column has been observed with trailing
    # padding and inconsistent case -- must still match after strip/lower.
    assert fix.guess_target_item_type("  NEWSPAPER ARTICLE   ", None) == "newspaperArticle"


def test_guess_target_item_type_magazine():
    assert fix.guess_target_item_type("Magazine Article", None) == "magazineArticle"


def test_guess_target_item_type_journal_article_is_left_alone():
    assert fix.guess_target_item_type("Journal Article", None) is None


def test_guess_target_item_type_falls_back_to_pub_type():
    assert fix.guess_target_item_type(None, "Newspaper Article") == "newspaperArticle"


def test_guess_target_item_type_unrelated_reference_type():
    assert fix.guess_target_item_type("Book", None) is None


def test_guess_target_item_type_empty():
    assert fix.guess_target_item_type(None, None) is None


# ---------------------------------------------------------------------------
# build_new_data
# ---------------------------------------------------------------------------

def _journal_article_data(**overrides):
    data = {
        "key": "ABC1",
        "version": 7,
        "itemType": "journalArticle",
        "title": "Decades of research burned in this Oregon forest",
        "creators": [{"creatorType": "author", "firstName": "Lynda V.", "lastName": "Mapes"}],
        "abstractNote": "It was a single lightning strike...",
        "publicationTitle": "The Seattle Times",
        "volume": "",
        "issue": "",
        "pages": "",
        "date": "2023",
        "url": "https://www.seattletimes.com/...",
        "ISSN": "",
        "tags": [{"tag": "pub_number:5327"}],
        "collections": ["COLLKEY1"],
        "relations": {},
        "extra": "Source DB: publication_id 5378; pub_number 5327; catalog_id 12092",
    }
    data.update(overrides)
    return data


def _newspaper_article_template():
    # Shape of zot.item_template("newspaperArticle")'s return value --
    # trimmed to the fields this module actually reads/writes plus a few
    # others, to catch accidental over-copying.
    return {
        "itemType": "newspaperArticle",
        "title": "",
        "creators": [],
        "abstractNote": "",
        "publicationTitle": "",
        "place": "",
        "edition": "",
        "date": "",
        "section": "",
        "pages": "",
        "language": "",
        "shortTitle": "",
        "url": "",
        "accessDate": "",
        "ISSN": "",
        "tags": [],
        "collections": [],
        "relations": {},
        "extra": "",
    }


def test_build_new_data_carries_over_shared_fields():
    old_data = _journal_article_data()
    template = _newspaper_article_template()

    new_data = fix.build_new_data(old_data, template, db_row={})

    assert new_data["title"] == old_data["title"]
    assert new_data["creators"] == old_data["creators"]
    assert new_data["publicationTitle"] == "The Seattle Times"
    assert new_data["date"] == "2023"
    assert new_data["url"] == old_data["url"]
    assert new_data["tags"] == old_data["tags"]
    assert new_data["collections"] == old_data["collections"]
    assert new_data["extra"] == old_data["extra"]


def test_build_new_data_drops_fields_with_no_home_in_new_type():
    old_data = _journal_article_data(volume="12", issue="3")
    template = _newspaper_article_template()

    new_data = fix.build_new_data(old_data, template, db_row={})

    # newspaperArticle has no volume/issue fields -- must not appear at all.
    assert "volume" not in new_data
    assert "issue" not in new_data


def test_build_new_data_preserves_key_and_version():
    old_data = _journal_article_data(key="XYZ9", version=42)
    template = _newspaper_article_template()

    new_data = fix.build_new_data(old_data, template, db_row={})

    assert new_data["key"] == "XYZ9"
    assert new_data["version"] == 42


def test_build_new_data_backfills_place_and_section_for_newspaper():
    old_data = _journal_article_data()
    template = _newspaper_article_template()
    db_row = {"place_published": "Seattle", "news_section": "Climate Lab"}

    new_data = fix.build_new_data(old_data, template, db_row)

    assert new_data["place"] == "Seattle"
    assert new_data["section"] == "Climate Lab"


def test_build_new_data_does_not_overwrite_existing_place():
    old_data = _journal_article_data()
    template = _newspaper_article_template()
    template["place"] = "already set"  # e.g. a human already filled this in

    new_data = fix.build_new_data(old_data, template, {"place_published": "Seattle"})

    assert new_data["place"] == "already set"


def test_build_new_data_no_place_section_backfill_for_magazine():
    old_data = _journal_article_data(itemType="journalArticle")
    magazine_template = {
        "itemType": "magazineArticle",
        "title": "",
        "creators": [],
        "publicationTitle": "",
        "volume": "",
        "issue": "",
        "date": "",
        "tags": [],
        "collections": [],
        "relations": {},
        "extra": "",
    }

    new_data = fix.build_new_data(
        old_data, magazine_template, {"place_published": "Seattle", "news_section": "Climate Lab"}
    )

    assert "place" not in new_data
    assert "section" not in new_data


# ---------------------------------------------------------------------------
# process_item
# ---------------------------------------------------------------------------

def _item(key, title="Some Title", extra="", item_type="journalArticle"):
    return {"data": {"key": key, "title": title, "itemType": item_type, "extra": extra}}


def _cfg(**overrides):
    base = dict(
        library_id="1",
        library_type="group",
        api_key="k",
        server="testserver",
        database="testdb",
        dry_run=True,
    )
    base.update(overrides)
    return fix.Config(**base)


def test_process_item_skips_no_pub_id():
    item = _item("A1", extra="no identifier here")
    row = fix.process_item(zot=MagicMock(), item=item, publication_id=None, db_row=None, cfg=_cfg())
    assert row.status == "skipped-no-pub-id"
    assert row.item_key == "A1"


def test_process_item_skips_not_in_db():
    item = _item("A2")
    row = fix.process_item(zot=MagicMock(), item=item, publication_id="9999", db_row=None, cfg=_cfg())
    assert row.status == "skipped-not-in-db"
    assert row.publication_id == "9999"


def test_process_item_skips_when_reference_type_is_really_journal():
    item = _item("A3")
    db_row = {"reference_type": "Journal Article", "pub_type": "Journal Article"}
    row = fix.process_item(zot=MagicMock(), item=item, publication_id="1", db_row=db_row, cfg=_cfg())
    assert row.status == "skipped-no-change-needed"


def test_process_item_dry_run_does_not_call_update():
    zot = MagicMock()
    item = _item("A4")
    db_row = {"reference_type": "Newspaper Article", "pub_type": "Newspaper Article"}

    row = fix.process_item(zot=zot, item=item, publication_id="5378", db_row=db_row, cfg=_cfg(dry_run=True))

    assert row.status == "dry-run"
    assert row.new_item_type == "newspaperArticle"
    zot.item.assert_not_called()
    zot.update_item.assert_not_called()


def test_process_item_live_converts_and_preserves_other_fields():
    fresh = {
        "data": {
            "key": "A5",
            "version": 42,
            "itemType": "journalArticle",
            "title": "Decades of research burned in this Oregon forest",
            "creators": [{"creatorType": "author", "lastName": "Mapes"}],
            "publicationTitle": "The Seattle Times",
            "tags": [{"tag": "pub_number:5327"}],
            "extra": "Source DB: publication_id 5378; pub_number 5327; catalog_id 12092",
        }
    }
    zot = MagicMock()
    zot.item.return_value = fresh
    zot.item_template.return_value = _newspaper_article_template()
    item = _item("A5")
    db_row = {
        "reference_type": "Newspaper Article",
        "pub_type": "Newspaper Article",
        "place_published": "Seattle",
        "news_section": "Climate Lab",
    }

    row = fix.process_item(zot=zot, item=item, publication_id="5378", db_row=db_row, cfg=_cfg(dry_run=False))

    assert row.status == "updated"
    zot.item.assert_called_once_with("A5")
    zot.item_template.assert_called_once_with("newspaperArticle")
    updated_item = zot.update_item.call_args[0][0]
    assert updated_item["data"]["itemType"] == "newspaperArticle"
    assert updated_item["data"]["place"] == "Seattle"
    assert updated_item["data"]["section"] == "Climate Lab"
    # fields shared with the old type survive
    assert updated_item["data"]["creators"] == [{"creatorType": "author", "lastName": "Mapes"}]
    assert updated_item["data"]["tags"] == [{"tag": "pub_number:5327"}]
    assert updated_item["data"]["key"] == "A5"
    assert updated_item["data"]["version"] == 42


def test_process_item_live_update_error_is_captured_not_raised():
    zot = MagicMock()
    zot.item.return_value = {"data": {"key": "A6", "itemType": "journalArticle"}}
    zot.item_template.return_value = _newspaper_article_template()
    zot.update_item.side_effect = RuntimeError("version conflict")
    item = _item("A6")
    db_row = {"reference_type": "Newspaper Article"}

    row = fix.process_item(zot=zot, item=item, publication_id="1", db_row=db_row, cfg=_cfg(dry_run=False))

    assert row.status == "error"
    assert "version conflict" in row.detail


# ---------------------------------------------------------------------------
# find_candidate_items
# ---------------------------------------------------------------------------

def test_find_candidate_items_filters_to_journal_article_and_paginates_top():
    zot = MagicMock()
    journal_item = _item("A", item_type="journalArticle")
    book_item = _item("B", item_type="book")
    zot.everything.return_value = [journal_item, book_item]

    result = fix.find_candidate_items(zot, collection_key=None)

    assert result == [journal_item]
    zot.top.assert_called_once()
    zot.collection_items.assert_not_called()


def test_find_candidate_items_uses_collection_when_given():
    zot = MagicMock()
    zot.everything.return_value = []

    fix.find_candidate_items(zot, collection_key="COLLKEY")

    zot.collection_items.assert_called_once_with("COLLKEY")
    zot.top.assert_not_called()


# ---------------------------------------------------------------------------
# find_pub_ids_by_parent_note
# ---------------------------------------------------------------------------

def _note(parent, text):
    return {"data": {"itemType": "note", "parentItem": parent, "note": text}}


def test_find_pub_ids_by_parent_note_extracts_from_html_note_body():
    zot = MagicMock()
    zot.everything.return_value = [
        _note("A", "<p>Notes: LTER8-3; media; Source DB: publication_id 5378; pub_number 5327; catalog_id 12092</p>"),
        _note("B", "<p>no identifier in this one</p>"),
    ]

    result = fix.find_pub_ids_by_parent_note(zot, collection_key=None)

    assert result == {"A": "5378"}
    zot.items.assert_called_once_with(itemType="note")
    zot.collection_items.assert_not_called()


def test_find_pub_ids_by_parent_note_uses_collection_when_given():
    zot = MagicMock()
    zot.everything.return_value = []

    fix.find_pub_ids_by_parent_note(zot, collection_key="COLLKEY")

    zot.collection_items.assert_called_once_with("COLLKEY", itemType="note")
    zot.items.assert_not_called()


def test_find_pub_ids_by_parent_note_skips_notes_with_no_parent():
    zot = MagicMock()
    zot.everything.return_value = [
        {"data": {"itemType": "note", "note": "Source DB: publication_id 1; pub_number 1; catalog_id 1"}},
    ]

    result = fix.find_pub_ids_by_parent_note(zot, collection_key=None)

    assert result == {}


# ---------------------------------------------------------------------------
# Missing optional dependencies
# ---------------------------------------------------------------------------

def test_get_client_raises_when_pyzotero_missing(monkeypatch):
    monkeypatch.setattr(fix, "zotero", None)
    cfg = _cfg()
    with pytest.raises(fix.ZoteroDepsMissingError):
        fix.get_client(cfg)


def test_fetch_reference_types_raises_when_pyodbc_missing(monkeypatch):
    monkeypatch.setattr(fix, "pyodbc", None)
    cfg = _cfg()
    with pytest.raises(fix.PyodbcMissingError):
        fix.fetch_reference_types(cfg, {"1"})


def test_fetch_reference_types_empty_ids_short_circuits():
    # Must not attempt a connection at all when there's nothing to look up.
    assert fix.fetch_reference_types(_cfg(), set()) == {}


def test_write_audit_xlsx_raises_when_openpyxl_missing(monkeypatch, tmp_path):
    monkeypatch.setattr(fix, "Workbook", None)
    with pytest.raises(fix.ZoteroDepsMissingError):
        fix.write_audit_xlsx([], str(tmp_path / "audit.xlsx"))


# ---------------------------------------------------------------------------
# load_config
# ---------------------------------------------------------------------------

def test_load_config_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        fix.load_config(tmp_path / "does_not_exist.ini")


def test_load_config_reads_zotero_and_database_sections(tmp_path):
    config_path = tmp_path / "db_config.ini"
    cp = configparser.ConfigParser()
    cp["Zotero"] = {"library_id": "12345", "library_type": "group", "api_key": "secret-key"}
    cp["Database"] = {
        "server": "testserver",
        "database": "testdb",
        "username": "testuser",
        "password": "testpwd",
        "trustservercertificate": "yes",
    }
    cp["Driver"] = {"driver": "{ODBC Driver 18 for SQL Server}"}
    with open(config_path, "w") as f:
        cp.write(f)

    result = fix.load_config(config_path)

    assert result == {
        "library_id": "12345",
        "library_type": "group",
        "api_key": "secret-key",
        "server": "testserver",
        "database": "testdb",
        "uid": "testuser",
        "pwd": "testpwd",
        "driver": "ODBC Driver 18 for SQL Server",
        "trust_server_certificate": True,
    }


# ---------------------------------------------------------------------------
# write_audit_xlsx
# ---------------------------------------------------------------------------

def test_write_audit_xlsx(tmp_path):
    rows = [
        fix.AuditRow("A", "Title A", "journalArticle", "newspaperArticle", "5378", "updated"),
        fix.AuditRow("B", "Title B", "journalArticle", "", "9999", "skipped-not-in-db"),
    ]
    out = tmp_path / "audit.xlsx"

    fix.write_audit_xlsx(rows, str(out))

    from openpyxl import load_workbook

    ws = load_workbook(out).active
    header = [c.value for c in ws[1]]
    assert header == fix.AUDIT_HEADERS
    assert [c.value for c in ws[2]][:5] == ["A", "Title A", "journalArticle", "newspaperArticle", "5378"]
    assert ws[2][5].value == "updated"
    assert ws[3][5].value == "skipped-not-in-db"


# ---------------------------------------------------------------------------
# run() orchestration
# ---------------------------------------------------------------------------

def test_run_raises_when_pyzotero_missing(monkeypatch):
    monkeypatch.setattr(fix, "zotero", None)
    cfg = _cfg(audit_path="unused.xlsx")
    with pytest.raises(fix.ZoteroDepsMissingError):
        fix.run(cfg, progress_callback=lambda *_: None)


def test_run_writes_audit_xlsx_and_counts(monkeypatch, tmp_path):
    fake_zotero_module = MagicMock()
    fake_client = MagicMock()
    # everything() just paginates whatever it's handed -- top()/items() are
    # mocked separately below since find_candidate_items() and
    # find_pub_ids_by_parent_note() each call it with a different source.
    fake_client.everything.side_effect = lambda source: source
    fake_client.top.return_value = [
        _item("A"),
        _item("B", item_type="book"),  # not a candidate at all
        _item("C"),  # candidate but unmatchable -- no linked note
    ]
    fake_client.items.return_value = [
        _note("A", "<p>Notes: media; Source DB: publication_id 5378; pub_number 5327; catalog_id 12092</p>"),
    ]
    fake_zotero_module.Zotero.return_value = fake_client
    monkeypatch.setattr(fix, "zotero", fake_zotero_module)
    monkeypatch.setattr(
        fix,
        "fetch_reference_types",
        lambda cfg, ids: {"5378": {"reference_type": "Newspaper Article"}} if ids else {},
    )

    audit_path = tmp_path / "audit.xlsx"
    cfg = _cfg(audit_path=str(audit_path))
    messages = []

    result = fix.run(cfg, progress_callback=messages.append)

    assert result.counts == {"dry-run": 1, "skipped-no-pub-id": 1}
    assert audit_path.exists()
    assert any("Found 2 Journal Article item(s)" in m for m in messages)
