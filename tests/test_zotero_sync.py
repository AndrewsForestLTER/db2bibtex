"""Tests for db2bibtex.zotero_sync: type/field mapping, matching, planning,
and writing. pyzotero, pyodbc and the DOI lookups are mocked throughout."""

from __future__ import annotations

from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from db2bibtex import zotero_sync as sync

# Trimmed versions of real Zotero API item templates.
TEMPLATE_FIELDS = {
    "journalArticle": ["title", "abstractNote", "publicationTitle", "volume", "issue", "pages",
                       "date", "DOI", "ISSN", "url", "extra"],
    "newspaperArticle": ["title", "abstractNote", "publicationTitle", "place", "date", "section",
                         "pages", "ISSN", "url", "extra"],
    "magazineArticle": ["title", "abstractNote", "publicationTitle", "volume", "issue", "date",
                        "pages", "ISSN", "url", "extra"],
    "bookSection": ["title", "abstractNote", "bookTitle", "edition", "place", "publisher", "date",
                    "pages", "ISBN", "url", "extra"],
    "book": ["title", "abstractNote", "series", "edition", "place", "publisher", "date", "ISBN",
             "url", "extra"],
    "report": ["title", "abstractNote", "seriesTitle", "place", "institution", "date", "pages",
               "url", "extra"],
    "thesis": ["title", "abstractNote", "thesisType", "university", "place", "date", "url", "extra"],
    "document": ["title", "abstractNote", "publisher", "date", "url", "extra"],
}


def _template(item_type, linkmode=None):
    if item_type == "note":
        return {"itemType": "note", "note": "", "tags": [], "collections": [], "relations": {}}
    if item_type == "attachment":
        return {"itemType": "attachment", "linkMode": linkmode, "title": "", "url": "",
                "contentType": "", "tags": [], "relations": {}}
    tmpl = {"itemType": item_type, "creators": [], "tags": [], "collections": [], "relations": {}}
    tmpl.update({f: "" for f in TEMPLATE_FIELDS[item_type]})
    return tmpl


def _row(pid, **overrides):
    row = {
        "publication_id": pid,
        "catalog_id": 9000 + int(pid),
        "pub_number": 100 + int(pid),
        "reference_type": "Journal Article",
        "pub_type": "Article",
        "author": "Doe, John//Doe, Jane",
        "pub_year": "2020",
        "title": f"Title {pid}",
        "secondary_author": None,
        "secondary_title": "Ecology",
        "place_published": None,
        "publisher": None,
        "volume": "10",
        "issue": "2",
        "pages": "1-10",
        "doi": f"10.1000/{pid}",
        "news_section": None,
        "edition": None,
        "type_of_work": None,
        "isbn_issn": None,
        "notes": None,
        "online_linkage": None,
        "online_pdf": None,
        "pdf": "F",
        "abstract": None,
    }
    row.update(overrides)
    return row


def _parent(key, item_type="journalArticle", **fields):
    data = _template(item_type)
    data.update({"key": key, "version": 5, "dateAdded": f"2026-01-01T00:00:{key[-1]}0Z"})
    data.update(fields)
    return {"data": data}


def _note(parent, pid):
    return {"data": {"key": f"N{parent}", "itemType": "note", "parentItem": parent,
                     "note": f"<p>Source DB: publication_id {pid}; pub_number 1; catalog_id 1</p>"}}


def _synced_parent(key, pid, **overrides):
    """A Zotero item that already matches _row(pid) exactly."""
    fields = {
        "title": f"Title {pid}",
        "creators": [
            {"creatorType": "author", "lastName": "Doe", "firstName": "John"},
            {"creatorType": "author", "lastName": "Doe", "firstName": "Jane"},
        ],
        "date": "2020",
        "publicationTitle": "Ecology",
        "volume": "10",
        "issue": "2",
        "pages": "1-10",
        "DOI": f"10.1000/{pid}",
        "tags": [{"tag": f"pub_number:{100 + int(pid)}"}],
    }
    fields.update(overrides)
    return _parent(key, **fields)


def _client(items):
    zot = MagicMock()
    zot.items.side_effect = lambda limit, start, **kw: items[start:start + limit]
    zot.item_template.side_effect = _template
    return zot


def _cfg(**overrides):
    cfg = sync.Config(library_id="1", library_type="group", api_key="k", server="s", database="d")
    for k, v in overrides.items():
        setattr(cfg, k, v)
    return cfg


def _plan(items, rows, **cfg_overrides):
    return sync.build_plan(_client(items), rows, _cfg(**cfg_overrides), lambda *_: None)


def _by_pid(plan):
    return {a.publication_id: a for a in plan.actions}


# ---------------------------------------------------------------------------
# Mapping
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "reference_type,expected",
    [
        ("Journal Article", "journalArticle"),
        ("  NEWSPAPER ARTICLE ", "newspaperArticle"),
        ("Magazine Article", "magazineArticle"),
        ("Book Section", "bookSection"),
        ("Edited Book", "book"),
        ("Thesis", "thesis"),
        ("Government Document", "report"),
        ("Something Else", "document"),
    ],
)
def test_guess_zotero_type(reference_type, expected):
    assert sync.guess_zotero_type(reference_type, None) == expected


def test_creator_from_author_two_field_and_single_field():
    assert sync.creator_from_author("Doe, John") == {"creatorType": "author", "lastName": "Doe", "firstName": "John"}
    assert sync.creator_from_author("USDA Forest Service") == {"creatorType": "author", "name": "USDA Forest Service"}


def test_creators_key_treats_two_field_and_single_field_alike():
    a = [{"creatorType": "author", "lastName": "Doe", "firstName": "John"}]
    b = [{"creatorType": "author", "name": "John  Doe"}]
    assert sync.creators_key(a) == sync.creators_key(b)


def test_db_fields_journal_article():
    out = sync.db_fields(_row("1"), _template("journalArticle"))
    assert out["publicationTitle"] == "Ecology"
    assert out["DOI"] == "10.1000/1"
    assert out["date"] == "2020"
    assert len(out["creators"]) == 2


def test_db_fields_places_secondary_title_and_publisher_by_type():
    section = sync.db_fields(_row("1", secondary_title="Big Book", publisher="Wiley"), _template("bookSection"))
    assert section["bookTitle"] == "Big Book" and section["publisher"] == "Wiley"
    report = sync.db_fields(_row("1", secondary_title="PNW-GTR", publisher="USFS"), _template("report"))
    assert report["seriesTitle"] == "PNW-GTR" and report["institution"] == "USFS"
    thesis = sync.db_fields(_row("1", publisher="OSU", type_of_work="M.S. thesis"), _template("thesis"))
    assert thesis["university"] == "OSU" and thesis["thesisType"] == "M.S. thesis"


def test_db_fields_doi_goes_to_extra_when_type_has_no_doi_field():
    out = sync.db_fields(_row("1"), _template("report"))
    assert "DOI" not in out and out["_extra_doi"] == "10.1000/1"


def test_db_fields_editors_only_for_editor_types():
    row = _row("1", secondary_author="Smith, Ann")
    assert [c["creatorType"] for c in sync.db_fields(row, _template("bookSection"))["creators"]] == ["author", "author", "editor"]
    assert all(c["creatorType"] == "author" for c in sync.db_fields(row, _template("journalArticle"))["creators"])


def test_db_fields_pdf_and_url():
    out = sync.db_fields(_row("1", pdf="T"), _template("journalArticle"))
    assert out["_pdf"].endswith("/pubs/pdf/pub101.pdf")
    assert out["url"].endswith("/publications/101")


def test_apply_fields_empty_db_value_never_blanks_zotero():
    data = _template("journalArticle")
    data["abstractNote"] = "Typed in by hand"
    wanted = sync.db_fields(_row("1", abstract=None), _template("journalArticle"))
    sync.apply_fields(data, wanted, 101)
    assert data["abstractNote"] == "Typed in by hand"


def test_apply_fields_merges_extra_doi_and_keeps_other_lines():
    data = _template("report")
    data["extra"] = "Original note line"
    changed = sync.apply_fields(data, {"_extra_doi": "10.1/x", "creators": []}, None)
    assert data["extra"] == "DOI: 10.1/x\nOriginal note line"
    assert changed == ["extra (DOI)"]


def test_norm_doi_strips_resolver():
    assert sync.norm_doi("https://doi.org/10.1000/ABC") == "10.1000/abc"
    assert sync.norm_doi("doi: 10.1000/abc") == "10.1000/abc"


# ---------------------------------------------------------------------------
# Planning
# ---------------------------------------------------------------------------

def test_plan_create_update_unchanged():
    items = [
        _synced_parent("AAAA1", "1"),
        _note("AAAA1", "1"),
        _synced_parent("BBBB2", "2", volume="99"),
        _note("BBBB2", "2"),
    ]
    plan = _plan(items, [_row("1"), _row("2"), _row("3")])
    acts = _by_pid(plan)
    assert acts["1"].status == "unchanged"
    assert acts["2"].status == "update" and acts["2"].changes == ["volume"]
    assert acts["2"].data["volume"] == "10" and acts["2"].data["version"] == 5
    assert acts["3"].status == "create" and acts["3"].data["itemType"] == "journalArticle"
    assert "publication_id 3" in acts["3"].note_text


def test_plan_converts_item_type_and_keeps_collections():
    items = [
        _synced_parent("AAAA1", "1", collections=["COLL"]),
        _note("AAAA1", "1"),
    ]
    plan = _plan(items, [_row("1", reference_type="Newspaper Article", place_published="Seattle")])
    act = _by_pid(plan)["1"]
    assert act.status == "update"
    assert act.changes[0] == "itemType"
    assert act.data["itemType"] == "newspaperArticle"
    assert act.data["place"] == "Seattle"
    assert act.data["collections"] == ["COLL"]
    assert "volume" not in act.data  # no home in newspaperArticle


def test_plan_adopts_unlinked_item_by_doi():
    items = [_synced_parent("ZZZZ9", "4", title="Different wording")]
    plan = _plan(items, [_row("4")])
    act = _by_pid(plan)["4"]
    assert act.status == "adopt" and act.item_key == "ZZZZ9"
    assert act.note_text and "title" in act.changes
    assert not plan.zotero_only


def test_plan_adopts_unlinked_item_by_title_and_year():
    items = [_parent("ZZZZ9", title="Title 4!", date="2020-05-01")]
    plan = _plan(items, [_row("4", doi=None)])
    assert _by_pid(plan)["4"].status == "adopt"


def test_plan_unlinked_duplicate_of_linked_item_is_reported_not_adopted():
    items = [_synced_parent("AAAA1", "1"), _note("AAAA1", "1"), _synced_parent("ZZZZ9", "1")]
    plan = _plan(items, [_row("1")])
    assert _by_pid(plan)["1"].status == "unchanged"
    assert [z.item_key for z in plan.zotero_only] == ["ZZZZ9"]
    assert "duplicate of publication_id 1" in plan.zotero_only[0].reason


def test_plan_unmatched_zotero_item_is_kept():
    plan = _plan([_parent("ZZZZ9", title="Brand new 2026 paper", date="2026")], [_row("1")])
    assert plan.zotero_only[0].reason.startswith("no matching database record")
    assert _by_pid(plan)["1"].status == "create"


def test_plan_ambiguous_match_is_not_adopted():
    items = [_parent("ZZZZ9", title="Same", date="2020")]
    plan = _plan(items, [_row("1", title="Same", doi=None), _row("2", title="Same", doi=None)])
    assert "several database records" in plan.zotero_only[0].reason
    assert {a.status for a in plan.actions} == {"create"}


def test_plan_reports_not_in_query_and_zotero_duplicates():
    items = [
        _synced_parent("AAAA1", "1"),
        _note("AAAA1", "1"),
        _synced_parent("AAAA2", "1"),
        _note("AAAA2", "1"),
        _synced_parent("CCCC3", "77"),
        _note("CCCC3", "77"),
    ]
    plan = _plan(items, [_row("1")])
    statuses = sorted((a.status, a.item_key) for a in plan.actions)
    assert ("zotero-duplicate", "AAAA2") in statuses
    assert ("not-in-query", "CCCC3") in statuses
    assert ("unchanged", "AAAA1") in statuses


def test_plan_reports_db_duplicates():
    plan = _plan([], [_row("1", doi="10.1/same"), _row("2", doi="10.1/same", title="Other")])
    assert plan.db_duplicates[0][:3] == ("DOI", "10.1/same", ["1", "2"])


def test_plan_skips_rows_without_title_or_author():
    plan = _plan([], [_row("1", author=None)])
    assert plan.skipped_rows == ["1"] and not plan.actions


def test_plan_adds_missing_pdf_attachment_only():
    items = [_synced_parent("AAAA1", "1", url="https://andrewsforest.oregonstate.edu/publications/101"),
             _note("AAAA1", "1")]
    plan = _plan(items, [_row("1", pdf="T")])
    act = _by_pid(plan)["1"]
    assert act.status == "update" and act.changes == ["PDF attachment"]

    items.append({"data": {"key": "ATT1", "itemType": "attachment", "parentItem": "AAAA1",
                           "url": act.pdf_url}})
    assert _by_pid(_plan(items, [_row("1", pdf="T")]))["1"].status == "unchanged"


# ---------------------------------------------------------------------------
# "et al." creators
# ---------------------------------------------------------------------------

ET_AL_AUTHORS = "Doe, John//et al."


def test_et_al_keeps_complete_zotero_creators(monkeypatch):
    lookup = MagicMock()
    monkeypatch.setattr(sync.authors_mod, "fetch_authors_crossref", lookup)
    items = [_synced_parent("AAAA1", "1"), _note("AAAA1", "1")]
    act = _by_pid(_plan(items, [_row("1", author=ET_AL_AUTHORS)]))["1"]
    assert act.status == "unchanged"
    lookup.assert_not_called()


def test_et_al_looks_up_doi_for_new_items(monkeypatch):
    monkeypatch.setattr(
        sync.authors_mod, "fetch_authors_crossref",
        lambda doi, mailto: [{"given": "John", "family": "Doe"}, {"given": "Ann", "family": "Roe"}],
    )
    act = _by_pid(_plan([], [_row("1", author=ET_AL_AUTHORS)], crossref_rate_limit_sec=0))["1"]
    assert [c["lastName"] for c in act.data["creators"]] == ["Doe", "Roe"]
    assert act.detail == "full author list from DOI lookup"


def test_et_al_no_lookup_when_disabled(monkeypatch):
    lookup = MagicMock()
    monkeypatch.setattr(sync.authors_mod, "fetch_authors_crossref", lookup)
    _plan([], [_row("1", author=ET_AL_AUTHORS)], fix_authors=False)
    lookup.assert_not_called()


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------

def test_apply_plan_creates_parents_then_children():
    zot = _client([])
    plan = sync.build_plan(zot, [_row("1", pdf="T"), _row("2")], _cfg(), lambda *_: None)
    zot.create_items.side_effect = [
        {"successful": {"0": {"key": "NEW1"}}, "failed": {"1": {"code": 400, "message": "bad"}}},
        {"successful": {"0": {"key": "C1"}, "1": {"key": "C2"}}},
    ]
    sync.apply_plan(zot, plan, _cfg(dry_run=False), lambda *_: None)

    acts = _by_pid(plan)
    assert acts["1"].status == "created" and acts["1"].item_key == "NEW1"
    assert acts["2"].status == "error" and "400 bad" in acts["2"].detail
    children = zot.create_items.call_args_list[1].args[0]
    assert [c["itemType"] for c in children] == ["note", "attachment"]
    assert all(c["parentItem"] == "NEW1" for c in children)
    assert children[1]["linkMode"] == "linked_url"


def test_apply_plan_respects_limit():
    zot = _client([])
    plan = sync.build_plan(zot, [_row("1"), _row("2")], _cfg(), lambda *_: None)
    zot.create_items.side_effect = [{"successful": {"0": {"key": "NEW1"}}}, {"successful": {"0": {"key": "C1"}}}]
    sync.apply_plan(zot, plan, _cfg(dry_run=False, limit=1), lambda *_: None)
    assert len(zot.create_items.call_args_list[0].args[0]) == 1
    assert sorted(a.status for a in plan.actions) == ["create", "created"]


def test_apply_plan_batches_of_50():
    zot = _client([])
    rows = [_row(str(i)) for i in range(1, 121)]
    plan = sync.build_plan(zot, rows, _cfg(), lambda *_: None)
    zot.create_items.side_effect = lambda chunk: {"successful": {str(i): {"key": f"K{i}"} for i in range(len(chunk))}}
    sync.apply_plan(zot, plan, _cfg(dry_run=False), lambda *_: None)
    sizes = [len(c.args[0]) for c in zot.create_items.call_args_list]
    assert sizes[:3] == [50, 50, 20]


def test_run_dry_run_writes_audit_and_nothing_to_zotero(monkeypatch, tmp_path):
    zot = _client([_parent("ZZZZ9", title="Kept", date="2026")])
    fake_module = MagicMock()
    fake_module.Zotero.return_value = zot
    monkeypatch.setattr(sync, "zotero", fake_module)
    monkeypatch.setattr(sync, "fetch_rows", lambda cfg: [_row("1"), _row("2", doi="10.1000/1")])

    audit = tmp_path / "audit.xlsx"
    plan = sync.run(_cfg(audit_path=str(audit)), progress_callback=lambda *_: None)

    zot.create_items.assert_not_called()
    assert plan.counts() == {"create": 2, "zotero-only": 1}
    wb = load_workbook(audit)
    assert wb.sheetnames == ["actions", "zotero_only", "db_duplicates", "summary"]
    assert wb["db_duplicates"].max_row == 2
    assert wb["actions"].max_row == 3


def test_run_raises_when_pyzotero_missing(monkeypatch):
    monkeypatch.setattr(sync, "zotero", None)
    with pytest.raises(sync.ZoteroDepsMissingError):
        sync.run(_cfg(), progress_callback=lambda *_: None)


def test_scan_library_pages_by_offset_and_ignores_orphans(monkeypatch):
    monkeypatch.setattr(sync, "PAGE_SIZE", 2)
    zot = MagicMock()
    pages = {
        0: [_synced_parent("AAAA1", "1"), _note("AAAA1", "1")],
        2: [_note("GONE1", "2")],  # GONE1's parent is in the trash
    }
    zot.items.side_effect = lambda limit, start, **kw: pages[start]
    lib = sync.scan_library(zot)
    assert list(lib.parents) == ["AAAA1"]
    assert lib.pub_ids_by_parent == {"AAAA1": ["1"], "GONE1": ["2"]}
    assert [c.kwargs["start"] for c in zot.items.call_args_list] == [0, 2]


def test_scan_library_retries_transient_errors(monkeypatch):
    monkeypatch.setattr(sync, "RETRY_BASE_SEC", 0)
    zot = MagicMock()
    zot.items.side_effect = [RuntimeError("Code: 502"), [_synced_parent("AAAA1", "1")]]
    messages = []
    lib = sync.scan_library(zot, messages.append)
    assert list(lib.parents) == ["AAAA1"]
    assert any("retrying" in m for m in messages)


def test_scan_library_gives_up_after_retries(monkeypatch):
    monkeypatch.setattr(sync, "RETRY_BASE_SEC", 0)
    zot = MagicMock()
    zot.items.side_effect = RuntimeError("Code: 502")
    with pytest.raises(RuntimeError):
        sync.scan_library(zot)
    assert zot.items.call_count == sync.PAGE_RETRIES + 1


def test_orphan_notes_do_not_link():
    items = [_note("GONE1", "2")]
    assert _by_pid(_plan(items, [_row("2")]))["2"].status == "create"


def test_pdf_attachment_title_is_file_name():
    zot = _client([])
    plan = sync.build_plan(zot, [_row("1", pdf="T")], _cfg(), lambda *_: None)
    zot.create_items.side_effect = [{"successful": {"0": {"key": "NEW1"}}},
                                    {"successful": {"0": {"key": "C1"}, "1": {"key": "C2"}}}]
    sync.apply_plan(zot, plan, _cfg(dry_run=False), lambda *_: None)
    assert zot.create_items.call_args_list[1].args[0][1]["title"] == "pub101.pdf"


def test_whitespace_only_differences_are_unchanged_and_raw_text_is_written():
    abstract = 'First paragraph.\n\nSecond paragraph.'
    items = [_synced_parent('AAAA1', '1', abstractNote='First paragraph. Second   paragraph.'),
             _note('AAAA1', '1')]
    assert _by_pid(_plan(items, [_row('1', abstract=abstract)]))['1'].status == 'unchanged'
    created = _by_pid(_plan([], [_row('1', abstract=abstract)]))['1']
    assert created.data['abstractNote'] == abstract


def test_dashes_latex_and_isbn_hyphens_are_not_changes():
    items = [_synced_parent("AAAA1", "1", title="Radiation budget \u2013 a review",
                            publicationTitle="Society & Natural Resources", ISSN="1234-5678"),
             _note("AAAA1", "1")]
    row = _row("1", title="Radiation budget -- a review", secondary_title="Society \\& Natural Resources",
               isbn_issn="12345678")
    assert _by_pid(_plan(items, [row]))["1"].status == "unchanged"


def test_new_items_get_en_dash_and_unescaped_text():
    act = _by_pid(_plan([], [_row("1", title="A--B", secondary_title="X \\& Y")]))["1"]
    assert act.data["title"] == "A\u2013B"
    assert act.data["publicationTitle"] == "X & Y"


def test_urls_are_written_verbatim():
    act = _by_pid(_plan([], [_row("1", online_linkage="https://example.org/a--b")]))["1"]
    assert act.data["url"] == "https://example.org/a--b"


def test_item_merged_from_two_live_records_is_left_alone():
    item = _synced_parent("AAAA1", "1")
    items = [item, _note("AAAA1", "1"),
             {"data": {"key": "NX", "itemType": "note", "parentItem": "AAAA1",
                       "note": "<p>Source DB: publication_id 2; pub_number 2; catalog_id 2</p>"}}]
    plan = _plan(items, [_row("1"), _row("2", volume="77")])
    acts = _by_pid(plan)
    assert acts["1"].status == acts["2"].status == "merged-duplicate"
    assert "publication_id 2" in acts["1"].detail
    assert not any(a.status == "create" for a in plan.actions)


def test_item_merged_with_out_of_query_record_syncs_normally():
    items = [_synced_parent("AAAA1", "1", volume="99"), _note("AAAA1", "1"),
             {"data": {"key": "NX", "itemType": "note", "parentItem": "AAAA1",
                       "note": "<p>Source DB: publication_id 2; pub_number 2; catalog_id 2</p>"}}]
    acts = _by_pid(_plan(items, [_row("1")]))
    assert acts["1"].status == "update"
    assert acts["2"].status == "not-in-query"


def test_triple_hyphen_is_em_dash():
    assert sync.clean_text("yr---1 and 1--2") == "yr—1 and 1–2"
